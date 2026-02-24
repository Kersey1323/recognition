from docx import Document
from docx.shared import Inches

doc = Document()
doc.add_heading('图像特征自动化提取流水线架构', 0)
doc.add_paragraph('以下为 SAM 零样本切割与 DINOv2 / ResNet-34 双引擎特征提取的完整多模态流水线图解：')
doc.add_picture('diagram.png', width=Inches(6.0))

output_path = 'docs/图像特征自动提取算法流程图.docx'
doc.save(output_path)
print(f'Document saved to {output_path}')

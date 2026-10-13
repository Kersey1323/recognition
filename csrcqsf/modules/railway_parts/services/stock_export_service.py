"""导出指定库区、货架中尚未采集视频的配件取件清单。"""

import io
import re

from anyio.to_thread import run_sync
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from csrcqsf.infra.db.main_db import SessionLocal
from csrcqsf.models.orm.item import Items
from csrcqsf.modules.railway_parts.schemas.parts import ExportUncollectedItemsReq


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LOCATION_PATTERN = re.compile(r"[^0-9]*([0-9]{2})([0-9]{2})([0-9]{2,3})")


def prepare_export_rows(items, req: ExportUncollectedItemsReq) -> list[tuple[str, str]]:
    records = []
    for item in items:
        positions = {}
        for location in item.stock_locations or []:
            if location.get("LGORT") != req.LGORT:
                continue
            code = str(location.get("GRARD") or "").strip()
            match = LOCATION_PATTERN.fullmatch(code)
            if match is None or match[1] != req.rack_no:
                continue
            positions[code] = (int(match[2]), int(match[3]), code)
        if not positions:
            continue
        ordered = sorted(positions, key=positions.get)
        # 同一配件只占一个名额；仅导出本次库区与货架内的位置，按最前面的排排序。
        key = (*positions[ordered[0]], item.item_name, int(item.id))
        records.append((key, (item.item_name, "\n".join(ordered))))
    records.sort(key=lambda record: record[0])
    return [row for _key, row in records[:req.limit]]


def build_workbook(rows: list[tuple[str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "未采集配件"
    sheet.append(("配件名称", "货位号"))
    for row in rows:
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", value) for value in row])
        for cell in sheet[sheet.max_row]:
            # 名称按文字写入，避免以等号开头的配件名称被 Excel 当作公式。
            cell.data_type = "s"
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        name_lines = sum(max(1, (len(line) + 21) // 22) for line in row[0].splitlines())
        sheet.row_dimensions[sheet.max_row].height = max(24, 16 * max(name_lines, len(row[1].splitlines())))
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=11, bold=True)
        cell.fill = PatternFill("solid", fgColor="E6EEF2")
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 32
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{sheet.max_row}"
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


async def export_uncollected_items(req: ExportUncollectedItemsReq) -> bytes:
    async with SessionLocal() as session:
        items = await Items.list_uncollected_stock_items(session, LGORT=req.LGORT, rack_no=req.rack_no)
    rows = prepare_export_rows(items, req)
    return await run_sync(build_workbook, rows)

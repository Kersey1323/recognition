import io
from io import StringIO

import pandas as pd
from anyio.to_thread import run_sync
from bfrcommon.snowflake import get_id

from csrcqsf.core.logger import CodeTimer
from csrcqsf.core.settings import minio_bucket
from csrcqsf.infra.db.main_db import SessionLocal
from csrcqsf.infra.external.basis_api import grpc_client
from csrcqsf.infra.storage.minio import minio_instance
from csrcqsf.models.orm.files import Files, FileTables
from globals_vars import logger


class DocBiz:
    @classmethod
    def get_paragraph_before_table(cls, doc, table_index=0, start_element_index=0):
        if table_index < 0 or table_index >= len(doc.tables):
            print(f"表格索引 {table_index} 无效。文档中共有 {len(doc.tables)} 个表格。")
            return "", 0

        target_table = doc.tables[table_index]
        previous_paragraph_text = ""
        loop_count = 0
        for index, element in enumerate(doc.element.body[start_element_index:], start_element_index):
            loop_count += 1
            if element == target_table._element:
                logger.info(f"为了找到表名遍历次数为{loop_count}，索引index{index}")
                return previous_paragraph_text, index
            if element.tag.endswith("p"):
                paragraph = None
                for p in doc.paragraphs:
                    if p._element == element:
                        paragraph = p
                        break
                if paragraph:
                    previous_paragraph_text = paragraph.text.strip()
        return "", 0

    @classmethod
    def rename_duplicate_columns(cls, df: pd.DataFrame):
        cols = pd.Series(df.columns)
        for dup in cols[cols.duplicated()].unique():
            indices = cols[cols == dup].index
            suffixes = [f"_{i + 1}" if i > 0 else "" for i in range(len(indices))]
            cols.iloc[indices] = [dup + suffix for suffix in suffixes]
        df.columns = cols
        return df

    @classmethod
    async def create_file(cls, file_name: str, file_url: str):
        json_data_dict: dict[str, str] = {}
        file_name_arr = file_url.split("/")
        if len(file_name_arr) < 2:
            raise Exception("文件名格式不正确")
        aim_file_name = file_name_arr[1]
        response = None

        try:
            response = await run_sync(minio_instance.client.get_object, minio_bucket, aim_file_name)
            if response.status != 200:
                raise Exception("oss请求出错")
            now_element_index = 0
            if response.data:
                try:
                    from docx import Document

                    docs = Document(io.BytesIO(response.data))
                    for index, table in enumerate(docs.tables):
                        table_data = []
                        logger.info(f"开始处理第{index + 1}个表，总共有{len(docs.tables)},进度{index / len(docs.tables) * 100}%")
                        for row in table.rows:
                            row_data = [cell.text.strip().replace(" ", "") for cell in row.cells]
                            table_data.append(row_data)
                        if table_data:
                            header = "".join(table_data[0])
                            if "图号" in header and "物料编码" in header:
                                df = pd.DataFrame(table_data[1:], columns=table_data[0])
                                if len(df.columns[df.columns.duplicated()]) > 0:
                                    df = cls.rename_duplicate_columns(df)
                                with CodeTimer("找到表格名"):
                                    table_name, now_element_index = await run_sync(
                                        cls.get_paragraph_before_table, docs, index, now_element_index
                                    )
                                json_data_dict[table_name] = df.to_json(
                                    force_ascii=False, orient="records", indent=4
                                )
                except Exception as e:
                    logger.error(f"文档解析出错{e}")
                    raise Exception("必须为docx格式的文档")
        finally:
            if response:
                response.close()
                response.release_conn()

        if json_data_dict:
            tables = []
            async with SessionLocal() as session:
                data = Files(file_name=file_name, file_url=file_url, table_num=len(json_data_dict), id=get_id())
                session.add(data)
                for table_name, json_data in json_data_dict.items():
                    tables.append(
                        FileTables(
                            file_id=data.id,
                            table_name=table_name,
                            table_data=json_data,
                        )
                    )
                if tables:
                    session.add_all(tables)
                await session.commit()

    @classmethod
    async def list_files(cls, page: int, per_page: int, query: str = None):
        async with SessionLocal() as session:
            return await Files.list_files(session, query, page, per_page)

    @classmethod
    async def list_tables(cls, file_id: int, page: int, per_page: int, query: str):
        async with SessionLocal() as session:
            return await FileTables.list_tables(session, file_id, query, page, per_page)

    @classmethod
    async def write_to_excel(cls, file_id: int, user_id: int):
        async with SessionLocal() as session:
            file = await Files.get_file(session, file_id)
            if file is None:
                raise Exception("源文件不存在")

        tables = await cls.list_tables(file_id, -1, 0, None)
        obj_name = f"{get_id()}.xlsx"
        excel_buffer = io.BytesIO()
        import xlsxwriter  # noqa: F401

        with pd.ExcelWriter(excel_buffer, engine="xlsxwriter") as writer:
            for table in tables.data:
                df = pd.read_json(StringIO(table.table_data))
                df.to_excel(writer, sheet_name=table.table_name, index=False)

        excel_buffer.seek(0)
        minio_instance.client.put_object(
            minio_bucket,
            obj_name,
            excel_buffer,
            length=excel_buffer.getbuffer().nbytes,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        grpc_client.send_file_task(
            obj_name,
            excel_buffer.getbuffer().nbytes,
            user_id,
            file.file_name,
        )

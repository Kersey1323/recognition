from docx import Document
from docx.shared import Inches

doc = Document()
doc.add_heading('SAM 大模型自动化图像切割流水线', 0)
doc.add_paragraph('以下特制为了详细展示 Segment Anything Model (SAM) 单独进行零样本抠图清洗及图形学背景白化的核心链路：')
doc.add_picture('diagram_sam_only.png', width=Inches(6.0))

output_path = 'docs/SAM自动化图像提取流水线_V2.docx'
doc.save(output_path)
print(f'Document saved to {output_path}')

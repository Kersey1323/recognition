import asyncio
import importlib
import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import pytest
import httpx
from fastapi import FastAPI
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.dialects import mysql

from csrcqsf.models.orm.item import Items
from csrcqsf.modules.railway_parts.schemas import _parts_api_models as schemas


def request(**kwargs):
    model = getattr(schemas, "ExportUncollectedItemsReq", None)
    assert model is not None, "export request model has not been added"
    return model(LGORT=" JS00 ", rack_no="9", **kwargs)


@pytest.mark.parametrize("value", [None, "", " ", "20", 20])
def test_export_limit_defaults_and_string_values(value):
    req = request(limit=value)
    assert req.limit == 20
    assert req.LGORT == "JS00"
    assert req.rack_no == "09"


def test_export_omitted_limit_defaults_to_twenty():
    assert request().limit == 20


@pytest.mark.parametrize("value", [0, -1, "abc", True])
def test_export_rejects_invalid_limit(value):
    assert getattr(schemas, "ExportUncollectedItemsReq", None) is not None
    with pytest.raises(ValidationError):
        request(limit=value)


def export_module():
    name = "csrcqsf.modules.railway_parts.services.stock_export_service"
    assert importlib.util.find_spec(name) is not None, "stock export service has not been added"
    return importlib.import_module(name)


def location(code, warehouse="JS00"):
    return {"LGORT": warehouse, "GRARD": code}


def test_export_sorts_rows_before_limit_and_excludes_other_locations():
    service = export_module()
    rows = [
        SimpleNamespace(id=3, item_name="第三排", stock_locations=[location("DF090301")]),
        SimpleNamespace(id=2, item_name="第二排", stock_locations=[location("DF090201")]),
        SimpleNamespace(id=1, item_name="第一排", stock_locations=[
            location("DF0901002"), location("DF090101"), location("DF090101"),
            location("DF100101"), location("DF090001", "2514"),
        ]),
    ]
    result = service.prepare_export_rows(rows, request(limit=2))
    assert result == [("第一排", "DF090101\nDF0901002"), ("第二排", "DF090201")]


def test_workbook_has_two_text_columns_and_supports_empty_results():
    service = export_module()
    data = service.build_workbook([("=配件名称", "DF090101\nDF0901002")])
    book = load_workbook(io.BytesIO(data))
    sheet = book.active
    assert tuple(cell.value for cell in sheet[1]) == ("配件名称", "货位号")
    assert sheet["A2"].value == "=配件名称"
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].value == "DF090101\nDF0901002"
    assert sheet.freeze_panes == "A2"
    empty = load_workbook(io.BytesIO(service.build_workbook([])))
    assert empty.active.max_row == 1


def test_export_repository_only_reads_uncollected_names_and_locations():
    method = getattr(Items, "list_uncollected_stock_items", None)
    assert method is not None, "export repository query has not been added"
    session = AsyncMock()
    session.execute.return_value = SimpleNamespace(all=lambda: [])
    asyncio.run(method(session, LGORT="JS00", rack_no="09"))
    compiled = session.execute.await_args.args[0].compile(dialect=mysql.dialect())
    query = str(compiled)
    assert "source_video_path IS NULL" in query
    assert "json_table(" in query
    assert "JS00" in compiled.params.values()
    assert "^[^0-9]*09[0-9]{4,5}$" in compiled.params.values()
    assert "ORDER BY" not in query
    assert "LIMIT" not in query


def test_export_route_returns_a_downloadable_excel_and_validates_input():
    service = export_module()
    routes = importlib.import_module("csrcqsf.modules.railway_parts.routes.parts_routes")
    app = FastAPI()
    app.include_router(routes.parts_router)
    payload = service.build_workbook([("配件一", "DF090101")])
    build = AsyncMock(return_value=payload)

    async def run():
        with patch.object(routes, "build_uncollected_items_export", build):
            async with httpx.AsyncClient(
                transport=httpx.ASGITransport(app=app), base_url="http://test",
            ) as client:
                response = await client.post("/railway/parts/export_uncollected_items", json={
                    "LGORT": "JS00", "rack_no": "09", "limit": "",
                })
                assert response.status_code == 200
                assert response.headers["content-type"] == service.XLSX_MEDIA_TYPE
                assert "attachment;" in response.headers["content-disposition"]
                assert response.content == payload
                invalid = await client.post("/railway/parts/export_uncollected_items", json={
                    "LGORT": "JS00", "rack_no": "09", "limit": 0,
                })
                assert invalid.status_code == 422
    asyncio.run(run())
    assert build.await_count == 1
    assert build.await_args.args[0].limit == 20
